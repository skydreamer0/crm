"""
週行程 Excel 讀取器 (weekly_plan.py)
=====================================

從「Weekly plan_<姓名>_MMDD.xlsx」取出指定日期的待訪名單，供排程自動執行使用。

工作表版面 (W1~W5 每張一週):
  - A 欄「日期」那列: B~H 為週一~週日 (新表為日期格式，舊表為「9月28日」文字)
  - A 欄「上午」/「下午」那列: 該列本身是表頭 (醫院/科別/客戶/等級08:30)，
    或是 國定假日/例假日/會議/颱風假 等標記 (常以合併儲存格蓋住整天) → 該時段跳過
  - 表頭下方直到下一個 A 欄標籤為止: 每格一筆 `醫院/科別/客戶/等級`
  - 「Night call」以下為週計畫，不屬於名單

非 4 段格式的格子 (例如「博田/採購（追蹤ACE500排裝）」) 不送 CRM，列入 skipped。
"""
# === 標準庫 ===
import datetime
import re
from dataclasses import dataclass, field
from pathlib import Path

# === 第三方套件 ===
import openpyxl

PLAN_FILE_PATTERN = re.compile(r"^weekly plan.*\.xls[xm]$", re.IGNORECASE)
PERIOD_LABELS = ("上午", "下午")
DATE_LABEL = "日期"
TEMPLATE_HINT = "醫院/科別"
TEXT_DATE_PATTERN = re.compile(r"^\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*$")
ENTRY_SPLIT_PATTERN = re.compile(r"[/／]")


@dataclass
class DayPlan:
    """One day's visit list extracted from the weekly plan."""
    date: datetime.date
    source: Path
    sheet: str
    column: str
    # [(時段, 名單文字)]，上午在前
    entries: list[tuple[str, str]] = field(default_factory=list)
    # [(儲存格, 內容, 原因)]
    skipped: list[tuple[str, str, str]] = field(default_factory=list)
    # {時段: 標記} 例如 {"上午": "國定假日"}
    off_periods: dict[str, str] = field(default_factory=dict)

    @property
    def lines(self) -> list[str]:
        return [line for _, line in self.entries]


def find_plan_files(folder: Path) -> list[Path]:
    """
    List weekly plan files, best candidate first.

    匯出版 (檔名含 out) 排在工作版之後；同類依修改時間新到舊。
    """
    files = [
        p for p in Path(folder).iterdir()
        if p.is_file() and not p.name.startswith("~$") and PLAN_FILE_PATTERN.match(p.name)
    ]
    return sorted(files, key=lambda p: ("out" in p.stem.lower(), -p.stat().st_mtime))


def find_day_plan(target: datetime.date, folder: Path | None = None, plan_file: Path | None = None) -> DayPlan | None:
    """Return the plan for `target` from `plan_file`, or the first file in `folder` containing it."""
    candidates = [Path(plan_file)] if plan_file else find_plan_files(folder)
    for path in candidates:
        plan = read_day_plan(path, target)
        if plan is not None:
            return plan
    return None


def read_day_plan(path: Path, target: datetime.date) -> DayPlan | None:
    """Read one workbook; return None when no sheet has a column for `target`."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        for ws in wb.worksheets:
            plan = _read_sheet(ws, Path(path), target)
            if plan is not None:
                return plan
    finally:
        wb.close()
    return None


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def _merged_value(ws, row: int, col: int):
    """Value of a cell, resolving merged ranges to their top-left cell."""
    cell = ws.cell(row, col)
    if cell.value is not None:
        return cell.value
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return None


def _to_date(value, target: datetime.date) -> datetime.date | None:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    match = TEXT_DATE_PATTERN.match(_cell_text(value))
    if not match:
        return None
    month, day = int(match.group(1)), int(match.group(2))
    # 舊表沒有年份: 取最接近目標日期的年份 (處理跨年週)
    year = target.year
    if month - target.month > 6:
        year -= 1
    elif target.month - month > 6:
        year += 1
    try:
        return datetime.date(year, month, day)
    except ValueError:
        return None


def _label_rows(ws) -> dict[str, int]:
    """Map column-A labels (日期/上午/下午/...) to their first row."""
    rows: dict[str, int] = {}
    for row in range(1, ws.max_row + 1):
        label = _cell_text(ws.cell(row, 1).value)
        if label and label not in rows:
            rows[label] = row
    return rows


def _period_end(ws, header_row: int) -> int:
    """Last row of a period block: stop at the next non-empty column-A label."""
    row = header_row + 1
    while row <= ws.max_row and not _cell_text(ws.cell(row, 1).value):
        row += 1
    return row - 1


def is_valid_entry(text: str) -> bool:
    """A sendable cell has exactly four non-empty `/`-separated parts."""
    parts = [p.strip() for p in ENTRY_SPLIT_PATTERN.split(text)]
    return len(parts) == 4 and all(parts) and TEMPLATE_HINT not in text


def _read_sheet(ws, path: Path, target: datetime.date) -> DayPlan | None:
    labels = _label_rows(ws)
    date_row = labels.get(DATE_LABEL)
    if date_row is None:
        return None

    target_col = None
    for col in range(2, ws.max_column + 1):
        if _to_date(ws.cell(date_row, col).value, target) == target:
            target_col = col
            break
    if target_col is None:
        return None

    column = openpyxl.utils.get_column_letter(target_col)
    plan = DayPlan(date=target, source=path, sheet=ws.title, column=column)

    for period in PERIOD_LABELS:
        header_row = labels.get(period)
        if header_row is None:
            continue
        header = _cell_text(_merged_value(ws, header_row, target_col))
        if header and TEMPLATE_HINT not in header:
            plan.off_periods[period] = header
            continue
        for row in range(header_row + 1, _period_end(ws, header_row) + 1):
            cell = ws.cell(row, target_col)
            text = _cell_text(cell.value)
            if not text:
                continue
            if is_valid_entry(text):
                plan.entries.append((period, text))
            else:
                plan.skipped.append((cell.coordinate, text, "非 醫院/科別/客戶/等級 四段格式"))

    return plan
