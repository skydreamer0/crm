"""
Tests for weekly_plan.py and scripts/run_today.py (synthetic workbooks only).
"""
import datetime
import importlib.util
import json
import logging
import os
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from weekly_plan import find_day_plan, find_plan_files, is_valid_entry, read_day_plan


def _load_run_today():
    cwd = os.getcwd()
    spec = importlib.util.spec_from_file_location("run_today", ROOT / "scripts" / "run_today.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    os.chdir(cwd)  # run_today chdirs to the project root on import
    return module


run_today = _load_run_today()


def _write_week(ws, dates, days):
    """
    Lay out one week like the real plan: row 2 dates, 上午 at row 3, 下午 at row 9.

    `days` maps column index (2=週一) to {"am": [...], "pm": [...], "off": "國定假日"}.
    """
    ws["A1"], ws["A2"], ws["A3"], ws["A9"], ws["A15"] = "星期", "日期", "上午", "下午", "Night call"
    ws.merge_cells("A3:A8")
    ws.merge_cells("A9:A14")
    for offset, value in enumerate(dates):
        ws.cell(2, 2 + offset, value)
    for col, day in days.items():
        if "off" in day:
            ws.cell(3, col, day["off"])
            ws.merge_cells(start_row=3, start_column=col, end_row=16, end_column=col)
            continue
        ws.cell(3, col, "醫院/科別/客戶/等級08:30")
        ws.cell(9, col, "醫院/科別/客戶/等級13:30")
        for i, line in enumerate(day.get("am", [])):
            ws.cell(4 + i, col, line)
        for i, line in enumerate(day.get("pm", [])):
            ws.cell(10 + i, col, line)
    ws["C15"] = "某院/採購/電話追蹤/A"
    ws["A17"] = "本週計畫(進藥、議價、活動安排、活動追蹤)"
    ws["A19"] = "某院/採購"


@pytest.fixture
def plan_path(tmp_path):
    wb = openpyxl.Workbook()
    old = wb.active
    old.title = "W1"
    _write_week(old, [f"9月{d}日" for d in range(28, 31)] + ["10月1日"], {
        2: {"off": "國定假日"},
        3: {"am": ["甲醫院/藥劑科/王小明/A"], "pm": ["乙醫院/採購/李小華/B"]},
    })
    new = wb.create_sheet("W5")
    dates = [datetime.datetime(2026, 10, 26) + datetime.timedelta(days=i) for i in range(7)]
    _write_week(new, dates, {
        2: {"off": "國定假日"},
        3: {
            "am": ["甲醫院/藥劑科/王小明/A", "甲醫院/感染科/陳大同/C"],
            "pm": ["乙醫院/採購/李小華/B", "丙醫院/採購（追蹤排裝）", "丙醫院/藥庫/張三/C"],
        },
        4: {"off": "會議"},
        7: {"off": "例假日"},
    })
    contacts = wb.create_sheet("聯絡人")
    contacts["A1"] = "客戶"
    path = tmp_path / "Weekly plan_Test_1023.xlsx"
    wb.save(path)
    return path


class TestReadDayPlan:

    def test_datetime_column_keeps_morning_before_afternoon(self, plan_path):
        plan = read_day_plan(plan_path, datetime.date(2026, 10, 27))
        assert plan.sheet == "W5" and plan.column == "C"
        assert plan.entries == [
            ("上午", "甲醫院/藥劑科/王小明/A"),
            ("上午", "甲醫院/感染科/陳大同/C"),
            ("下午", "乙醫院/採購/李小華/B"),
            ("下午", "丙醫院/藥庫/張三/C"),
        ]

    def test_non_four_part_cells_are_skipped_not_sent(self, plan_path):
        plan = read_day_plan(plan_path, datetime.date(2026, 10, 27))
        assert [(coord, text) for coord, text, _ in plan.skipped] == [("C11", "丙醫院/採購（追蹤排裝）")]

    @pytest.mark.parametrize("day, mark", [(26, "國定假日"), (28, "會議"), (31, "例假日")])
    def test_off_days_have_no_entries(self, plan_path, day, mark):
        plan = read_day_plan(plan_path, datetime.date(2026, 10, day))
        assert plan.entries == []
        assert plan.off_periods == {"上午": mark, "下午": mark}

    def test_text_dates_from_old_sheets(self, plan_path):
        plan = read_day_plan(plan_path, datetime.date(2026, 9, 29))
        assert plan.sheet == "W1"
        assert plan.lines == ["甲醫院/藥劑科/王小明/A", "乙醫院/採購/李小華/B"]

    def test_night_call_row_is_not_an_entry(self, plan_path):
        plan = read_day_plan(plan_path, datetime.date(2026, 10, 27))
        texts = [text for _, text in plan.entries] + [text for _, text, _ in plan.skipped]
        assert all("電話追蹤" not in text for text in texts)

    def test_date_outside_workbook_returns_none(self, plan_path):
        assert read_day_plan(plan_path, datetime.date(2026, 12, 1)) is None


class TestFindPlanFiles:

    def test_working_copy_preferred_over_out_and_lock_files(self, plan_path, tmp_path):
        out = tmp_path / "Weekly plan_Test_1023out.xlsx"
        out.write_bytes(plan_path.read_bytes())
        (tmp_path / "~$Weekly plan_Test_1023.xlsx").write_bytes(b"lock")
        (tmp_path / "notes.xlsx").write_bytes(b"")
        os.utime(out, (plan_path.stat().st_mtime + 60,) * 2)
        assert find_plan_files(tmp_path) == [plan_path, out]

    def test_find_day_plan_searches_folder(self, plan_path, tmp_path):
        plan = find_day_plan(datetime.date(2026, 10, 27), folder=tmp_path)
        assert plan.source == plan_path


@pytest.mark.parametrize("text, valid", [
    ("甲醫院/藥劑科/王小明/A", True),
    ("甲醫院／藥劑科／王小明／A", True),
    ("丙醫院/採購（追蹤排裝）", False),
    ("甲醫院/藥劑科/王小明", False),
    ("醫院/科別/客戶/等級08:30", False),
])
def test_is_valid_entry(text, valid):
    assert is_valid_entry(text) is valid


class TestRunToday:

    def test_build_entries_uses_hpk_parser_and_keeps_period(self, plan_path):
        plan = read_day_plan(plan_path, datetime.date(2026, 10, 27))
        entries = run_today.build_entries(plan)
        assert [(e.visit_period, e.hospital_name, e.customer_name) for e in entries] == [
            ("上午", "甲醫院", "王小明"),
            ("上午", "甲醫院", "陳大同"),
            ("下午", "乙醫院", "李小華"),
            ("下午", "丙醫院", "張三"),
        ]
        assert all(e.matched_products == [] for e in entries)

    @pytest.mark.parametrize("now, expected", [
        (datetime.datetime(2026, 10, 27, 20, 30), (datetime.date(2026, 10, 27), False)),
        (datetime.datetime(2026, 10, 28, 8, 0), (datetime.date(2026, 10, 27), True)),
    ])
    def test_run_date_before_cutoff_catches_up_previous_day(self, now, expected):
        assert run_today.choose_run_date(None, now) == expected

    def test_explicit_date_is_never_a_catch_up(self):
        now = datetime.datetime(2026, 10, 28, 8, 0)
        assert run_today.choose_run_date("2026-10-28", now) == (datetime.date(2026, 10, 28), False)

    def test_previous_runs_only_count_after_daily_report_step(self, tmp_path):
        def write(name, record):
            (tmp_path / name).write_text(json.dumps(record), encoding="utf-8")

        write("run_a.json", {"start_time": "2026-10-27T20:30:00", "logs": [{"step": "login"}, {"step": "error"}]})
        write("run_b.json", {"start_time": "2026-10-27T20:40:00", "logs": [{"step": "login"}, {"step": "daily_report"}]})
        write("run_c.json", {"start_time": "2026-10-27T21:00:00", "run_date": "2026-10-26",
                             "logs": [{"step": "daily_report"}]})
        found = run_today.find_previous_runs(datetime.date(2026, 10, 27), history_dir=tmp_path)
        assert [p.name for p in found] == ["run_b.json"]

    def test_browser_present_skips_install(self, tmp_path):
        exe = tmp_path / "chrome.exe"
        exe.write_bytes(b"")
        run_today.ensure_browser_installed(executable=lambda: exe, run=lambda *a, **k: pytest.fail("installed"))

    def test_missing_browser_is_installed_once(self, tmp_path):
        exe = tmp_path / "chrome.exe"
        calls = []

        def fake_run(cmd, **kwargs):
            calls.append(cmd)
            exe.write_bytes(b"")
            return type("Result", (), {"returncode": 0, "stdout": "", "stderr": ""})()

        run_today.ensure_browser_installed(executable=lambda: exe, run=fake_run)
        assert calls == [[sys.executable, "-m", "playwright", "install", "chromium"]]

    def test_failed_install_raises(self, tmp_path):
        exe = tmp_path / "chrome.exe"
        failed = lambda *a, **k: type("Result", (), {"returncode": 1, "stdout": "", "stderr": "network error"})()
        with pytest.raises(RuntimeError, match="network error"):
            run_today.ensure_browser_installed(executable=lambda: exe, run=failed)

    def test_dry_run_never_touches_crm(self, plan_path, tmp_path, monkeypatch):
        monkeypatch.setattr(run_today, "SCHEDULED_LOG_DIR", tmp_path / "logs")
        monkeypatch.setattr(run_today, "message_box", lambda *a, **k: pytest.fail("dialog shown"))
        monkeypatch.setitem(sys.modules, "create_appointments", None)  # import would fail
        root = logging.getLogger()
        handlers = list(root.handlers)
        try:
            code = run_today.main(["--date", "2026-10-27", "--plan-file", str(plan_path), "--dry-run"])
        finally:
            for handler in root.handlers[len(handlers):]:
                handler.close()
            root.handlers = handlers
        assert code == run_today.EXIT_OK
        log_text = (tmp_path / "logs" / "run_20261027.log").read_text(encoding="utf-8")
        assert "王小明" in log_text and "dry-run" in log_text
